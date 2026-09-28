import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_user_journey_vertical_flowchart(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    slide = prs.slides.add_slide(blank_layout)

    # 1. Clean White Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg.line.fill.background()

    # Color Palette for White Background
    COLOR_TITLE = RGBColor(15, 23, 42)      # Slate 900
    COLOR_SUBTITLE = RGBColor(71, 85, 105)  # Slate 600

    # Step Card Themes (Step 1 to 5)
    ACCENTS = [
        {"color": RGBColor(2, 132, 199),   "bg": RGBColor(240, 249, 255), "role": "USER ONBOARDING"},
        {"color": RGBColor(124, 58, 237),  "bg": RGBColor(245, 243, 255), "role": "REAL-TIME MONITORING"},
        {"color": RGBColor(217, 119, 6),   "bg": RGBColor(254, 243, 199), "role": "AI FORECASTING"},
        {"color": RGBColor(225, 29, 72),   "bg": RGBColor(255, 241, 242), "role": "GIS IMPACT ASSESSMENT"},
        {"color": RGBColor(16, 185, 129),  "bg": RGBColor(236, 253, 245), "role": "ACTION & DECISION"},
    ]

    # 2. Header Block
    tag = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.38), Inches(3.4), Inches(0.28))
    tag.fill.solid()
    tag.fill.fore_color.rgb = RGBColor(238, 242, 255)
    tag.line.color.rgb = RGBColor(199, 210, 254)
    tag.line.width = Pt(1)
    tf_tag = tag.text_frame
    tf_tag.vertical_anchor = MSO_ANCHOR.MIDDLE
    pt = tf_tag.paragraphs[0]
    pt.text = "●  HOW A USER INTERACTS WITH THE PROTOTYPE"
    pt.font.size = Pt(8.5)
    pt.font.bold = True
    pt.font.color.rgb = RGBColor(79, 70, 229)

    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.66), Inches(11.7), Inches(0.58))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = "User Experience Flowchart: How Users & Authorities Use VAYU-DRISHTI"
    p_title.font.size = Pt(18)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_TITLE

    p_sub = tf_title.add_paragraph()
    p_sub.text = "Step-by-step user journey from landing on the portal to viewing live AI predictions, assessing district risk, and executing scenario shifts."
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = COLOR_SUBTITLE

    # 3. User Journey Steps
    user_steps = [
        {
            "num": "STEP 1",
            "action": "ACCESS PORTAL & SELECT DATA MODE",
            "user_intent": "User chooses what type of cyclone intelligence to inspect",
            "what_user_does": [
                "Clicks prototype URL (instant load, no login required for public views)",
                "Toggles Data Mode: 'Live data' (real-time stream) or 'Historical' (Amphan, Fani, etc.)",
                "Chooses role perspective: 'Tactical Monitor' (Authorities) or 'Citizen View' (Public)"
            ],
            "prototype_response": "Loads dynamic cyclone catalog, initializes high-speed Carto Voyager basemap, and synchronizes real-time IST clock."
        },
        {
            "num": "STEP 2",
            "action": "TRACK LIVE STORM & INTERACTIVE TIMELINE",
            "user_intent": "User inspects storm position, structure, and historical progression",
            "what_user_does": [
                "Selects active cyclone from the dropdown list (e.g. Cyclone Vayu-Drishti / Amphan)",
                "Drags interactive Timeline Slider or clicks '▶ Play' to scrub through lifecycle",
                "Toggles Map Layers: Storm Eye Marker, Wind Radii, Observed Track, Places"
            ],
            "prototype_response": "Map automatically flies and fits bounds to storm coordinates; animates storm eye marker and displays real-time pressure & sustained wind metrics."
        },
        {
            "num": "STEP 3",
            "action": "ANALYZE AI MULTI-HORIZON PREDICTIONS",
            "user_intent": "User explores where the storm will head next and how strong it will get",
            "what_user_does": [
                "Examines the AI Forecast Trajectory (+6h, +12h, +18h, +24h, +36h, +48h)",
                "Inspects the expanding Cone of Uncertainty along the projected landfall path",
                "Checks Intensity Trend Badge: Strengthening, Stable, or Weakening",
                "Reviews explainability drivers: Sea Surface Temperature (SST) & Vertical Wind Shear"
            ],
            "prototype_response": "Executes kinematic persistence & deep neural forecast models, returning calibrated track coordinates with evaluated benchmark error bounds."
        },
        {
            "num": "STEP 4",
            "action": "ASSESS GIS DISTRICT & POPULATION EXPOSURE",
            "user_intent": "User determines which administrative areas & infrastructure are in danger",
            "what_user_does": [
                "Views highlighted Risk Corridor polygon overlaid directly onto Indian districts",
                "Scans the District Exposure Table (sorted by highest population risk & landfall threat)",
                "Inspects critical coastal infrastructure tallies: Major Ports, Airports, NH Highways, Hospitals"
            ],
            "prototype_response": "Runs spatial intersections across 700+ Indian districts via GeoPandas/Shapely and estimates exposed population from census density references."
        },
        {
            "num": "STEP 5",
            "action": "SIMULATE 'WHAT-IF' SHIFTS & TAKE ACTION",
            "user_intent": "Authorities run contingency scenarios; Citizens access safety guides",
            "what_user_does": [
                "AUTHORITY: Adjusts 'Shift Track' slider (± km West/East) & changes intensity scenario",
                "AUTHORITY: Reviews dynamic deltas (+ added / - dropped districts) & exports SITREP report",
                "CITIZEN: Selects local district to view nearest shelters, helplines, and 3-phase safety checklists"
            ],
            "prototype_response": "Recalculates transformed geometry and population exposure deltas in real-time (< 300ms) and provides bilingual citizen emergency directives."
        }
    ]

    card_left = Inches(0.8)
    card_width = Inches(11.7)
    card_height = Inches(0.92)
    step_gap = Inches(0.18)
    start_top = Inches(1.5)

    for i, s in enumerate(user_steps):
        top_y = start_top + i * (card_height + step_gap)
        accent = ACCENTS[i]

        # Main Card Box
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, top_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = accent["bg"]
        card.line.color.rgb = accent["color"]
        card.line.width = Pt(1.5)

        # Step Pill Badge (Left)
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left + Inches(0.15), top_y + Inches(0.12), Inches(0.9), Inches(0.68))
        badge.fill.solid()
        badge.fill.fore_color.rgb = accent["color"]
        badge.line.fill.background()
        tf_b = badge.text_frame
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        pb = tf_b.paragraphs[0]
        pb.alignment = PP_ALIGN.CENTER
        pb.text = s["num"]
        pb.font.size = Pt(11)
        pb.font.bold = True
        pb.font.color.rgb = RGBColor(255, 255, 255)

        # Column 1: User Action & Goal
        tb_col1 = slide.shapes.add_textbox(card_left + Inches(1.15), top_y + Inches(0.06), Inches(3.6), Inches(0.8))
        tf_c1 = tb_col1.text_frame
        tf_c1.word_wrap = True
        p_c1 = tf_c1.paragraphs[0]
        p_c1.text = s["action"]
        p_c1.font.size = Pt(10)
        p_c1.font.bold = True
        p_c1.font.color.rgb = COLOR_TITLE

        p_sub_c1 = tf_c1.add_paragraph()
        p_sub_c1.text = s["user_intent"]
        p_sub_c1.font.size = Pt(8.5)
        p_sub_c1.font.color.rgb = accent["color"]
        p_sub_c1.font.bold = True

        # Column 2: What the User Actually Does (User Actions)
        tb_col2 = slide.shapes.add_textbox(card_left + Inches(4.85), top_y + Inches(0.06), Inches(3.9), Inches(0.8))
        tf_c2 = tb_col2.text_frame
        tf_c2.word_wrap = True
        for idx, act in enumerate(s["what_user_does"]):
            p_act = tf_c2.paragraphs[0] if idx == 0 else tf_c2.add_paragraph()
            p_act.text = f"• {act}"
            p_act.font.size = Pt(8.0)
            p_act.font.color.rgb = COLOR_TITLE
            p_act.space_after = Pt(2)

        # Column 3: Prototype Response / Outcome
        tb_col3 = slide.shapes.add_textbox(card_left + Inches(8.85), top_y + Inches(0.06), Inches(2.75), Inches(0.8))
        tf_c3 = tb_col3.text_frame
        tf_c3.word_wrap = True
        p_out_lbl = tf_c3.paragraphs[0]
        p_out_lbl.text = "⚡ System Response:"
        p_out_lbl.font.size = Pt(8.0)
        p_out_lbl.font.bold = True
        p_out_lbl.font.color.rgb = accent["color"]

        p_out = tf_c3.add_paragraph()
        p_out.text = s["prototype_response"]
        p_out.font.size = Pt(8.0)
        p_out.font.color.rgb = COLOR_SUBTITLE

        # Downward Flow Arrow between steps
        if i < len(user_steps) - 1:
            arr_top = top_y + card_height + Inches(0.02)
            arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, card_left + Inches(0.5), arr_top, Inches(0.18), Inches(0.14))
            arr.fill.solid()
            arr.fill.fore_color.rgb = RGBColor(148, 163, 184) # Slate 400
            arr.line.fill.background()

    # Footer note
    tb_foot = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.3))
    tf_foot = tb_foot.text_frame
    p_foot = tf_foot.paragraphs[0]
    p_foot.text = "VAYU-DRISHTI User Journey — Designed for Dual-Persona Operations: High-Level Tactical Command (NDMA/Collectors) & Public Citizen Early Warning"
    p_foot.font.size = Pt(8)
    p_foot.font.color.rgb = RGBColor(148, 163, 184)

    prs.save(output_path)
    print(f"[SUCCESS] User-POV vertical flowchart saved to: {output_path}")

if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("d:/cyclone/VAYU_DRISHTI_Prototype_Workflow.pptx")
    create_user_journey_vertical_flowchart(out)
